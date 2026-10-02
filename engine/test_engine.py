import io
import base64
from pathlib import Path
import sys
sys.path.insert(0,str(Path(__file__).parent))
import pytest
from PIL import Image
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from main import dispatch

@pytest.fixture
def source(tmp_path):
    path = tmp_path/"source.pdf"
    b=io.BytesIO()
    c=canvas.Canvas(b,pagesize=(400,500))
    image=Image.effect_noise((1800,1800),80).convert("RGB")
    for i in range(3):
        c.setFont("Helvetica",18)
        c.drawString(30,450,f"Original secret page {i+1}")
        c.drawImage(ImageReader(image),30,180,180,180)
        c.showPage()
    c.save()
    w=PdfWriter(clone_from=PdfReader(b))
    w.add_outline_item("Second",1)
    with path.open("wb") as f:w.write(f)
    return path

def job(source,tmp_path,tool,**kw):
    output=tmp_path/(tool+".pdf")
    return dict(tool=tool,inputs=[str(source)],output=str(output),**kw)

def test_preview_replace_text_and_image(source,tmp_path):
    preview=dispatch(dict(tool="preview",inputs=[str(source)],page=0))
    assert preview["count"]==3 and preview["image"].startswith("data:image/png")
    text=next(o for o in preview["objects"] if o["kind"]=="text")
    image=next(o for o in preview["objects"] if o["kind"]=="image")
    req=job(source,tmp_path,"edit",operations=[
        dict(kind="text",objectId=text["id"],page=0,x=text["x"],y=text["y"],text="Replacement text",size=18),
        dict(kind="delete",objectId=image["id"],page=0)])
    dispatch(req)
    result=PdfReader(req["output"])
    assert "Replacement text" in result.pages[0].extract_text()
    assert "Original secret" not in result.pages[0].extract_text()
    assert "Original secret" in result.pages[1].extract_text()
    assert len(result.pages[0].images)==0

def test_real_redaction_removes_text(source,tmp_path):
    req=job(source,tmp_path,"edit",operations=[dict(kind="redact",page=0,x=0,y=0,width=400,height=100)])
    dispatch(req)
    r=PdfReader(req["output"])
    assert all(not p.extract_text().strip() for p in r.pages)
    assert b"secret" not in Path(req["output"]).read_bytes()

def test_form_creation_and_filling(source,tmp_path):
    req=job(source,tmp_path,"edit",operations=[dict(kind="field",page=0,x=30,y=100,width=180,height=30,text="Customer")])
    dispatch(req)
    fields=dispatch(dict(tool="form_fields",inputs=[req["output"]]))["fields"]
    assert fields[0]["name"]=="Customer"
    out=tmp_path/"filled.pdf"
    dispatch(dict(tool="fill_forms",inputs=[req["output"]],output=str(out),values={"Customer":"Ada Lovelace"}))
    assert PdfReader(out).get_fields()["Customer"]["/V"]=="Ada Lovelace"

@pytest.mark.parametrize("tool,count,options",[
    ("half",6,{}),("nup",2,{"amount":2}),("flip",3,{}),("resize",3,{"mode":"letter"}),
    ("header_footer",3,{"text":"Page {page} / {pages}"}),("bates",3,{"text":"CASE-","amount":10}),
    ("bookmarks",3,{"text":"2: Chapter two"}),("repair",3,{})])
def test_page_operations_open_and_have_expected_pages(source,tmp_path,tool,count,options):
    req=job(source,tmp_path,tool,**options)
    dispatch(req)
    r=PdfReader(req["output"])
    assert len(r.pages)==count
    if tool=="resize":assert float(r.pages[0].mediabox.width)==612
    if tool=="bates":assert "CASE-000010" in r.pages[0].extract_text()

@pytest.mark.parametrize("tool,options,counts",[
    ("split_bookmarks",{},[1,2]),("split_text",{"text":"page 3"},[2,1]),
    ("split_every",{"amount":2},[2,1]),("rename_text",{},[1,1,1])])
def test_split_boundaries(source,tmp_path,tool,options,counts):
    req=job(source,tmp_path,tool,**options);req["output"]=str(tmp_path/tool)
    result=dispatch(req)
    assert [len(PdfReader(p).pages) for p in result["outputs"]]==counts

def test_compression_reduces_image_pdf_and_preserves_text(source,tmp_path):
    req=job(source,tmp_path,"compress",mode="small")
    dispatch(req)
    assert Path(req["output"]).stat().st_size < source.stat().st_size
    assert "Original secret" in PdfReader(req["output"]).pages[0].extract_text()

@pytest.mark.parametrize("tool,extension",[("to_word","docx"),("to_excel","xlsx"),("to_powerpoint","pptx")])
def test_office_files_open_and_contain_editable_text(source,tmp_path,tool,extension):
    req=job(source,tmp_path,tool);req["output"]=str(tmp_path/("converted."+extension))
    dispatch(req)
    if tool=="to_word":
        from docx import Document
        assert any("Original secret" in p.text for p in Document(req["output"]).paragraphs)
    elif tool=="to_excel":
        from openpyxl import load_workbook
        assert load_workbook(req["output"]).active["A1"].value.startswith("Original")
    else:
        from pptx import Presentation
        prs=Presentation(req["output"])
        assert len(prs.slides)==3
        assert any("Original secret" in s.text for s in prs.slides[0].shapes if s.has_text_frame)

def test_certificate_signature_integrity(source,tmp_path):
    from cryptography import x509
    from cryptography.x509.oid import NameOID
    from cryptography.hazmat.primitives import hashes,serialization
    from cryptography.hazmat.primitives.asymmetric import rsa
    from cryptography.hazmat.primitives.serialization import pkcs12
    from datetime import datetime,timedelta,timezone
    from pyhanko.pdf_utils.reader import PdfFileReader
    from pyhanko.sign.validation import validate_pdf_signature
    from pyhanko_certvalidator import ValidationContext
    key=rsa.generate_private_key(public_exponent=65537,key_size=2048)
    name=x509.Name([x509.NameAttribute(NameOID.COMMON_NAME,"pdfmerge test")])
    cert=x509.CertificateBuilder().subject_name(name).issuer_name(name).public_key(key.public_key()).serial_number(123).not_valid_before(datetime.now(timezone.utc)-timedelta(days=1)).not_valid_after(datetime.now(timezone.utc)+timedelta(days=1)).sign(key,hashes.SHA256())
    pfx=tmp_path/"test.pfx"
    pfx.write_bytes(pkcs12.serialize_key_and_certificates(b"test",key,cert,None,serialization.BestAvailableEncryption(b"pass")))
    req=job(source,tmp_path,"digital_sign",certificate=str(pfx),password="pass")
    dispatch(req)
    with open(req["output"],"rb") as f:
        pdf=PdfFileReader(f)
        assert len(pdf.embedded_signatures)==1
        status=validate_pdf_signature(pdf.embedded_signatures[0],signer_validation_context=ValidationContext(allow_fetching=False))
        assert status.intact and status.valid

def test_never_overwrite_input(source,tmp_path):
    original=source.read_bytes()
    with pytest.raises(ValueError):
        dispatch(dict(tool="repair",inputs=[str(source)],output=str(source)))
    assert source.read_bytes()==original

def test_visual_signature_and_links_persist(source,tmp_path):
    import base64
    from PIL import ImageDraw
    sig=Image.new("RGBA",(300,100),(0,0,0,0))
    ImageDraw.Draw(sig).line([(20,80),(120,10),(80,70),(240,40)],fill="black",width=5)
    b=io.BytesIO();sig.save(b,"PNG")
    req=job(source,tmp_path,"edit",operations=[
        dict(kind="signature",page=0,x=40,y=100,width=180,height=60,data="data:image/png;base64,"+base64.b64encode(b.getvalue()).decode()),
        dict(kind="link",page=0,x=40,y=400,width=150,height=20,url="https://example.com")])
    dispatch(req)
    r=PdfReader(req["output"])
    assert len(r.pages[0].images)==2
    assert any(a.get_object().get("/A",{}).get("/URI")=="https://example.com" for a in r.pages[0]["/Annots"])


def test_uploaded_signature_background_choices(source,tmp_path):
    import base64
    from PIL import ImageDraw
    picture=Image.new("RGB",(200,90),"white")
    ImageDraw.Draw(picture).line((15,70,70,15,150,55),fill=(15,35,100),width=5)
    source_image=tmp_path/"handwriting.png"
    picture.save(source_image)
    original=source_image.read_bytes()
    kept=dispatch(dict(tool="signature_image",inputs=[str(source_image)],removeBackground=False))
    removed=dispatch(dict(tool="signature_image",inputs=[str(source_image)],removeBackground=True))
    keep_image=Image.open(io.BytesIO(base64.b64decode(kept["data"].split(",")[1])))
    clear_image=Image.open(io.BytesIO(base64.b64decode(removed["data"].split(",")[1])))
    assert keep_image.getpixel((0,0))[3]==255
    assert clear_image.getpixel((0,0))[3]==0
    assert clear_image.width<keep_image.width
    assert source_image.read_bytes()==original
    req=job(source,tmp_path,"edit",operations=[dict(kind="signature",page=0,x=30,y=50,width=190,height=70,data=removed["data"])])
    dispatch(req)
    assert PdfReader(req["output"]).pages[0].images


def blank_signature_document(tmp_path, name="signature-pages.pdf", pages=3):
    path = tmp_path / name
    c = canvas.Canvas(str(path), pagesize=(400, 500))
    for _ in range(pages):
        c.showPage()
    c.save()
    return path


def signature_image_data(size=(1200, 400)):
    b = io.BytesIO()
    Image.new("RGB", size, "#163266").save(b, "PNG")
    return "data:image/png;base64," + base64.b64encode(b.getvalue()).decode()


def test_signature_instances_export_final_rectangles_without_resampling(tmp_path):
    source = blank_signature_document(tmp_path)
    original = source.read_bytes()
    data = signature_image_data()
    # An instance removed on page 2 is absent from the final list; editing an
    # instance on page 3 does not change either instance retained on page 1.
    final_instances = [
        dict(id="first", kind="signature", page=0, x=20, y=40, width=180, height=60, data=data),
        dict(id="duplicate", kind="signature", page=0, x=220, y=310, width=90, height=30, data=data),
        dict(id="resized", kind="signature", page=2, x=60, y=120, width=270, height=90, data=data),
    ]
    req = job(source, tmp_path, "edit", operations=final_instances)
    dispatch(req)
    result = PdfReader(req["output"])
    for page_index, expected in ((0, final_instances[:2]), (1, []), (2, final_instances[2:])):
        preview = dispatch(dict(tool="preview", inputs=[req["output"]], page=page_index))
        images = [obj for obj in preview["objects"] if obj["kind"] == "image"]
        assert len(images) == len(expected)
        for image, rectangle in zip(images, expected):
            assert [image[key] for key in ("x", "y", "width", "height")] == pytest.approx(
                [rectangle[key] for key in ("x", "y", "width", "height")], abs=0.001)
        assert not result.pages[page_index].get("/Annots")
        for image in result.pages[page_index].images:
            assert image.image.size == (1200, 400)
    assert source.read_bytes() == original


def test_drawn_signature_exports_vector_strokes_at_resized_rectangle(tmp_path):
    from PIL import ImageChops
    from common import render
    source = blank_signature_document(tmp_path, pages=1)
    req = job(source, tmp_path, "edit", operations=[dict(
        kind="signature", page=0, x=50, y=70, width=200, height=100,
        data=signature_image_data(), color="#000000", signatureStrokeWidth=4,
        signatureViewBox=dict(width=100, height=50),
        signaturePaths=[[[10, 10], [90, 40]], [[50, 25]]],
    )])
    dispatch(req)
    result = PdfReader(req["output"])
    # The fallback PNG is used by the UI only: export has vector paths, no bitmap.
    assert len(result.pages[0].images) == 0
    raster = render(Path(req["output"]).read_bytes(), 0, 144).convert("RGB")
    ink_bounds = ImageChops.difference(raster, Image.new("RGB", raster.size, "white")).getbbox()
    # At 2 pixels/point the stroke geometry includes its round end caps.
    assert ink_bounds == pytest.approx((132, 172, 468, 308), abs=2)


@pytest.mark.parametrize("rotation", [0, 90, 180, 270])
def test_signature_export_matches_cropped_rotated_preview(tmp_path, rotation):
    from pypdf.generic import RectangleObject
    source = blank_signature_document(tmp_path, pages=1)
    w = PdfWriter(clone_from=PdfReader(source))
    w.pages[0].cropbox = RectangleObject([40, 60, 340, 460])
    w.pages[0].rotate(rotation)
    cropped = tmp_path / "cropped.pdf"
    with cropped.open("wb") as output:
        w.write(output)
    preview = dispatch(dict(tool="preview", inputs=[str(cropped)], page=0))
    width, height = (300, 400) if rotation in (0, 180) else (400, 300)
    assert (preview["width"], preview["height"]) == (width, height)
    rectangle = dict(x=width - 120, y=height - 50, width=120, height=40)
    req = job(cropped, tmp_path, "edit", operations=[
        dict(kind="signature", page=0, data=signature_image_data(), **rectangle)])
    dispatch(req)
    exported = dispatch(dict(tool="preview", inputs=[req["output"]], page=0))
    image = next(obj for obj in exported["objects"] if obj["kind"] == "image")
    assert (exported["width"], exported["height"]) == (width, height)
    assert [image[key] for key in rectangle] == pytest.approx(list(rectangle.values()), abs=0.001)


def test_uploaded_signature_keeps_original_resolution(tmp_path):
    picture = tmp_path / "large-signature.png"
    Image.new("RGB", (3600, 1200), "#173155").save(picture)
    prepared = dispatch(dict(tool="signature_image", inputs=[str(picture)], removeBackground=False))
    image = Image.open(io.BytesIO(base64.b64decode(prepared["data"].split(",", 1)[1])))
    assert (prepared["width"], prepared["height"]) == image.size == (3600, 1200)


@pytest.mark.parametrize("rectangle", [
    dict(x=-1, y=20, width=90, height=30),
    dict(x=320, y=20, width=90, height=30),
    dict(x=20, y=480, width=90, height=30),
    dict(x=20, y=20, width=0, height=30),
    dict(x=float("nan"), y=20, width=90, height=30),
])
def test_signature_export_rejects_invalid_page_bounds(tmp_path, rectangle):
    source = blank_signature_document(tmp_path, pages=1)
    req = job(source, tmp_path, "edit", operations=[
        dict(kind="signature", page=0, data=signature_image_data(), **rectangle)])
    with pytest.raises(ValueError, match="outside its page"):
        dispatch(req)
    assert not Path(req["output"]).exists()


def test_editor_find_replace_links_annotations_and_forms(source,tmp_path):
    req=job(source,tmp_path,"edit",operations=[
        dict(kind="find_replace",page=0,scope="all",find="secret",replace="public",matchCase=False),
        dict(kind="link",page=0,x=30,y=20,width=50,height=20,targetPage=2),
        dict(kind="highlight",page=0,x=30,y=32,width=100,height=22,color="#fff000"),
        dict(kind="multiline",page=0,x=30,y=90,width=150,height=55,text="Notes"),
        dict(kind="dropdown",page=0,x=30,y=150,width=150,height=30,text="Status",options="New\nDone"),
        dict(kind="radio",page=0,x=30,y=200,width=20,height=20,text="Approval",value="Yes",selected=True),
        dict(kind="radio",page=0,x=70,y=200,width=20,height=20,text="Approval",value="No"),
        dict(kind="whiteout",page=0,x=0,y=260,width=30,height=30),
        dict(kind="arrow",page=0,x=30,y=310,width=50,height=30)])
    dispatch(req)
    r=PdfReader(req["output"])
    assert all("secret" not in page.extract_text().lower() for page in r.pages)
    assert all("public" in page.extract_text().lower() for page in r.pages)
    fields=r.get_fields()
    assert {"Notes","Status","Approval"}.issubset(fields)
    assert fields["Notes"]["/Ff"] & 4096
    annotations=[a.get_object() for a in r.pages[0]["/Annots"]]
    assert any(a.get("/Subtype")=="/Highlight" for a in annotations)
    assert any(a.get("/Subtype")=="/Link" and "/Dest" in a for a in annotations)


def test_editor_replaces_existing_link(source,tmp_path):
    first=job(source,tmp_path,"edit",operations=[dict(kind="link",page=0,x=30,y=50,width=90,height=20,url="https://old.example")])
    dispatch(first)
    link=next(obj for obj in dispatch(dict(tool="preview",inputs=[first["output"]],page=0))["objects"] if obj["kind"]=="link")
    second=tmp_path/"changed-link.pdf"
    dispatch(dict(tool="edit",inputs=[first["output"]],output=str(second),operations=[
        dict(kind="link",page=0,annotationId=link["annotationId"],x=link["x"],y=link["y"],
             width=link["width"],height=link["height"],url="https://new.example")]))
    links=[a.get_object() for a in PdfReader(second).pages[0]["/Annots"] if a.get_object().get("/Subtype")=="/Link"]
    assert len(links)==1
    assert links[0]["/A"]["/URI"]=="https://new.example"

def test_compare_extract_and_images_conversion(source,tmp_path):
    req=job(source,tmp_path,"compare");req["inputs"].append(str(source))
    result=dispatch(req)
    assert "0 pages differ" in result["message"]
    assert len(PdfReader(req["output"]).pages)==3
    folder=tmp_path/"images"
    result=dispatch(dict(tool="to_images",inputs=[str(source)],output=str(folder),pages="2",amount=72,mode="png"))
    assert len(result["outputs"])==1
    out=tmp_path/"from-images.pdf"
    dispatch(dict(tool="images_pdf",inputs=result["outputs"],output=str(out)))
    assert len(PdfReader(out).pages)==1

def test_grayscale_flatten_deskew_and_alternate(source,tmp_path):
    for tool in ("grayscale","flatten","deskew"):
        req=job(source,tmp_path,tool);dispatch(req)
        assert len(PdfReader(req["output"]).pages)==3
    req=job(source,tmp_path,"alternate");req["inputs"]=[str(source),str(source)]
    dispatch(req)
    r=PdfReader(req["output"])
    assert len(r.pages)==6
    assert r.pages[0].extract_text()==r.pages[1].extract_text()


def test_drawn_signature_clips_strokes_to_final_rectangle(tmp_path):
    from PIL import ImageChops
    from common import render
    source = blank_signature_document(tmp_path, pages=1)
    req = job(source, tmp_path, "edit", operations=[dict(
        kind="signature", page=0, x=50, y=70, width=200, height=100,
        color="#000000", signatureStrokeWidth=4,
        signatureViewBox=dict(width=100, height=50),
        signaturePaths=[[[-500, 25], [500, 25]]],
    )])
    dispatch(req)
    raster = render(Path(req["output"]).read_bytes(), 0, 144).convert("RGB")
    bounds = ImageChops.difference(raster, Image.new("RGB", raster.size, "white")).getbbox()
    assert bounds == pytest.approx((100, 232, 500, 248), abs=1)


def test_export_after_removing_all_signatures_has_no_overlays(tmp_path):
    source = blank_signature_document(tmp_path, pages=2)
    req = job(source, tmp_path, "edit", operations=[])
    dispatch(req)
    output = PdfReader(req["output"])
    assert len(output.pages) == 2
    assert all(len(page.images) == 0 for page in output.pages)


@pytest.mark.parametrize("rotation", [15, 45, 90, 180, 270])
def test_rotated_signature_image_export_bounds(tmp_path, rotation):
    from PIL import ImageChops
    from common import render
    import math
    source = blank_signature_document(tmp_path, pages=1)
    rectangle = dict(x=100, y=150, width=120, height=40)
    req = job(source, tmp_path, "edit", operations=[dict(kind="signature", page=0,
        rotation=rotation, data=signature_image_data(), **rectangle)])
    dispatch(req)
    raster = render(Path(req["output"]).read_bytes(), 0, 144).convert("RGB")
    bounds = ImageChops.difference(raster, Image.new("RGB", raster.size, "white")).getbbox()
    angle = math.radians(rotation)
    w = abs(math.cos(angle))*120 + abs(math.sin(angle))*40
    h = abs(math.sin(angle))*120 + abs(math.cos(angle))*40
    assert bounds == pytest.approx(((160-w/2)*2, (170-h/2)*2, (160+w/2)*2, (170+h/2)*2), abs=2)
    assert PdfReader(req["output"]).pages[0].images[0].image.size == (1200, 400)


def test_rotated_drawn_signature_exports_clockwise_vectors(tmp_path):
    from PIL import ImageChops
    from common import render
    source = blank_signature_document(tmp_path, pages=1)
    req = job(source, tmp_path, "edit", operations=[dict(kind="signature", page=0,
        x=100, y=150, width=120, height=40, rotation=90, color="#000000",
        signatureViewBox=dict(width=120,height=40), signatureStrokeWidth=4,
        signaturePaths=[[[10,10],[110,10]]])])
    dispatch(req)
    raster = render(Path(req["output"]).read_bytes(), 0, 144).convert("RGB")
    bounds = ImageChops.difference(raster, Image.new("RGB", raster.size, "white")).getbbox()
    assert bounds == pytest.approx((336,236,344,444), abs=2)
    assert not PdfReader(req["output"]).pages[0].images


def test_signature_rotation_checks_visible_bounds_instead_of_unrotated_frame(tmp_path):
    source = blank_signature_document(tmp_path, pages=1)
    invalid = job(source, tmp_path, "edit", operations=[dict(kind="signature",page=0,
        x=100,y=0,width=200,height=40,rotation=90,data=signature_image_data())])
    with pytest.raises(ValueError,match="outside its page"):
        dispatch(invalid)
    valid = job(source, tmp_path, "edit", operations=[dict(kind="signature",page=0,
        x=-80,y=100,width=200,height=40,rotation=90,data=signature_image_data())])
    dispatch(valid)
    assert Path(valid["output"]).exists()
