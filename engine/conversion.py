import io
import json
import os
from pathlib import Path
import re
import tempfile
from PIL import Image
from common import *


def office_convert(req):
    program = find_program("office")
    source = Path(req["inputs"][0]).resolve()
    with tempfile.TemporaryDirectory(prefix="pdfmerge-office-") as temp:
        temp = Path(temp)
        profile = (temp / "profile").as_uri()
        command = [program, "-env:UserInstallation=" + profile, "--headless", "--nologo", "--nodefault", "--norestore"]
        fmt = "pdf"
        if req["tool"] == "pdfa":
            command += ["--infilter=draw_pdf_import"]
            fmt = 'pdf:draw_pdf_Export:{"SelectPdfVersion":{"type":"long","value":"2"}}'
        command += ["--convert-to", fmt, "--outdir", str(temp), str(source)]
        run_process(command, timeout=360)
        result = temp / (source.stem + ".pdf")
        if not result.is_file():
            raise ValueError("Office conversion produced no PDF. Check that the input opens correctly.")
        # Verify this is a readable document before publishing it.
        reader(result)
        return outcome([atomic_bytes(req["output"], result.read_bytes())],
                       "PDF/A-2b exported through LibreOffice." if req["tool"] == "pdfa" else "Converted to PDF.")


def to_office(req):
    import pdfplumber
    tool = req["tool"]
    source = req["inputs"][0]
    output = io.BytesIO()
    with pdfplumber.open(source) as pdf:
        if tool == "to_word":
            from docx import Document
            from docx.shared import Inches
            doc = Document()
            for i, page in enumerate(pdf.pages):
                if i:
                    doc.add_page_break()
                text = page.extract_text(layout=False) or ""
                if text.strip():
                    for line in text.splitlines():
                        doc.add_paragraph(line)
                    # Preserve illustrations; text remains editable.
                    for item in page.images:
                        bounds = (max(0,item["x0"]), max(0,item["top"]), min(page.width,item["x1"]), min(page.height,item["bottom"]))
                        if bounds[2] <= bounds[0] or bounds[3] <= bounds[1]:
                            continue
                        image = render(source, i, 100)
                        scale = 100/72
                        image = image.crop(tuple(int(v*scale) for v in bounds))
                        b = io.BytesIO()
                        image.save(b,"PNG")
                        b.seek(0)
                        doc.add_picture(b, width=Inches(min(6,(bounds[2]-bounds[0])/72)))
                else:
                    image = render(source, i, 130)
                    b = io.BytesIO()
                    image.save(b, "PNG")
                    b.seek(0)
                    doc.add_picture(b, width=Inches(6))
                    doc.add_paragraph("This page is a scan. Run OCR first to obtain editable text.")
            doc.save(output)
        elif tool == "to_excel":
            from openpyxl import Workbook
            wb = Workbook()
            wb.remove(wb.active)
            for i, page in enumerate(pdf.pages):
                ws = wb.create_sheet(f"Page {i+1}")
                tables = page.extract_tables()
                if tables:
                    for table in tables:
                        for row in table:
                            ws.append([str(cell or "") for cell in row])
                        ws.append([])
                else:
                    lines = (page.extract_text() or "No selectable text. Run OCR first.").splitlines()
                    for line in lines:
                        ws.append([line])
                # PDF text must remain data, never executable spreadsheet formulas.
                for row in ws:
                    for cell in row:
                        if isinstance(cell.value, str):
                            cell.data_type = "s"
                for column in ws.columns:
                    ws.column_dimensions[column[0].column_letter].width = 28
            wb.save(output)
        elif tool == "to_powerpoint":
            from pptx import Presentation
            from pptx.util import Pt
            from pptx.dml.color import RGBColor
            import pypdfium2 as pdfium
            data = normalize(source)
            with pdfium.PdfDocument(data) as doc:
                prs = Presentation()
                first = doc[0]
                sw, sh = first.get_size()
                first.close()
                prs.slide_width, prs.slide_height = Pt(sw), Pt(sh)
                for index in range(len(doc)):
                    page = doc[index]
                    pw, ph = page.get_size()
                    textpage = page.get_textpage()
                    objects = list(page.get_objects(textpage=textpage))
                    text_items = []
                    for obj in objects:
                        if obj.type == pdfium.raw.FPDF_PAGEOBJ_TEXT:
                            text_items.append((obj.extract(), obj.get_bounds(), obj.get_font_size()))
                    textpage.close()
                    for obj in objects:
                        if obj.type == pdfium.raw.FPDF_PAGEOBJ_TEXT:
                            page.remove_obj(obj)
                            obj.close()
                    page.gen_content()
                    bitmap = page.render(scale=1.4)
                    b = io.BytesIO()
                    bitmap.to_pil().save(b, "PNG")
                    b.seek(0)
                    bitmap.close()
                    slide = prs.slides.add_slide(prs.slide_layouts[6])
                    slide.shapes.add_picture(b, 0, 0, width=Pt(sw), height=Pt(sh))
                    for text, (left,bottom,right,top), size in text_items:
                        box = slide.shapes.add_textbox(Pt(left*sw/pw), Pt((ph-top)*sh/ph),
                                                      Pt(max(10,right-left+10)*sw/pw), Pt(max(size*1.6,top-bottom+8)*sh/ph))
                        frame = box.text_frame
                        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
                        frame.word_wrap = False
                        frame.text = text
                        for para in frame.paragraphs:
                            para.font.size = Pt(size*sh/ph)
                            para.font.color.rgb = RGBColor(0,0,0)
                    page.close()
                prs.save(output)
        else:
            raise ValueError("Unknown Office output.")
    return outcome([atomic_bytes(req["output"], output.getvalue())],
                   "Office file saved. Review layout and tables; PDF conversion is best effort.")


def image_conversion(req):
    tool = req["tool"]
    if tool == "images_pdf":
        images = []
        for path in req["inputs"]:
            image = Image.open(path)
            for frame in range(getattr(image, "n_frames", 1)):
                image.seek(frame)
                images.append(image.convert("RGB").copy())
        return outcome([atomic_bytes(req["output"], image_pdf(images, jpeg_quality=95))])
    folder = Path(req["output"])
    folder.mkdir(parents=True, exist_ok=True)
    source = req["inputs"][0]
    if tool == "extract_images":
        r = reader(source)
        pending = []
        for i, page in enumerate(r.pages):
            for j, image in enumerate(page.images):
                extension = Path(image.name).suffix or ".bin"
                pending.append((folder / f"page-{i+1:03}-image-{j+1:03}{extension}", image.data))
    else:
        r = reader(source)
        pending = []
        fmt = req.get("mode", "png")
        for i in ranges(req.get("pages", ""), len(r.pages)):
            image = render(source, i, max(72, min(300, float(req.get("amount", 144)))))
            buf = io.BytesIO()
            image.convert("RGB").save(buf, "JPEG" if fmt == "jpg" else "PNG")
            pending.append((folder/f"page-{i+1:03}.{fmt}", buf.getvalue()))
    if not pending:
        raise ValueError("No embedded images were found.")
    if any(path.exists() for path, _ in pending):
        raise ValueError("An output already exists. Choose an empty folder.")
    return outcome([atomic_bytes(p, data) for p,data in pending])


def scan(req):
    if sys.platform != "win32":
        raise ValueError("Scanner capture is available in the Windows app.")
    with tempfile.TemporaryDirectory(prefix="pdfmerge-scan-") as temp:
        target = Path(temp)/"scan.bmp"
        env = dict(os.environ, PDFMERGE_SCAN_PATH=str(target))
        script = "$d=New-Object -ComObject WIA.CommonDialog; $i=$d.ShowAcquireImage(); if ($null -ne $i) { $i.SaveFile($env:PDFMERGE_SCAN_PATH) }"
        run_process(["powershell.exe","-NoProfile","-STA","-Command",script],timeout=600,env=env)
        if not target.exists():
            raise ValueError("Scanning was cancelled or no scanner is available.")
        image = Image.open(target)
        return outcome([atomic_bytes(req["output"],image_pdf([image],jpeg_quality=95))])


def ocr_or_deskew(req):
    import cv2
    import numpy as np
    tool = req["tool"]
    source = req["inputs"][0]
    r = reader(source)
    images, sizes = [], []
    writer = PdfWriter()
    if tool == "ocr":
        import pytesseract
        pytesseract.pytesseract.tesseract_cmd = find_program("ocr")
        tessdata = resource_root() / "tesseract/tessdata"
        config = '--tessdata-dir "' + str(tessdata) + '"' if tessdata.is_dir() else ""
    for i,p in enumerate(r.pages):
        image = render(source, i, 200).convert("RGB")
        if tool == "deskew":
            gray = cv2.cvtColor(np.array(image), cv2.COLOR_RGB2GRAY)
            binary = cv2.threshold(gray,0,255,cv2.THRESH_BINARY_INV+cv2.THRESH_OTSU)[1]
            edges = cv2.Canny(binary,50,150)
            lines = cv2.HoughLinesP(edges,1,np.pi/1800,100,minLineLength=max(40,image.width//6),maxLineGap=20)
            angles = []
            if lines is not None:
                for line in lines[:,0]:
                    x1,y1,x2,y2 = line
                    angle = np.degrees(np.arctan2(y2-y1,x2-x1))
                    if abs(angle) < 12:
                        angles.append(angle)
            if angles:
                image = image.rotate(float(np.median(angles)),resample=Image.Resampling.BICUBIC,expand=False,fillcolor="white")
            images.append(image)
            sizes.append((float(p.mediabox.width),float(p.mediabox.height)))
        else:
            blob = pytesseract.image_to_pdf_or_hocr(image,extension="pdf",lang=req.get("language","eng"),config=config)
            writer.append(PdfReader(io.BytesIO(blob)))
    data = image_pdf(images,sizes,95) if tool == "deskew" else pdf_bytes(writer)
    return outcome([atomic_bytes(req["output"],data)],"Searchable OCR PDF saved." if tool == "ocr" else "Deskewed image pages saved.")
